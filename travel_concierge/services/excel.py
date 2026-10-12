import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from travel_concierge.storage import xlsx_path


def write_options(data):
    path = xlsx_path(data)
    pd.DataFrame(data["rows"]).to_excel(path, sheet_name="Options", index=False)
    _style(path)
    return path


def add_booking(data):
    """Append a Booking sheet (simulated) to the existing workbook."""
    path = xlsx_path(data)
    b = data["booking"]
    with pd.ExcelWriter(path, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
        pd.DataFrame([{"SIMULATED": "No real booking made", **b["option"], "Reference": b["reference"],
                       "Booked At": b["time"]}]).to_excel(w, sheet_name="Booking", index=False)
    _style(path)
    return path


def _style(path):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        for i, cell in enumerate(ws[1], 1):
            if "(INR)" in str(cell.value):
                for (c,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    c.number_format = '"₹"#,##0'
        ws.freeze_panes = "A2"
    wb.save(path)

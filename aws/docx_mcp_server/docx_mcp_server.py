import io, json, os
import boto3
from docx import Document
from mcp.server.fastmcp import FastMCP

REGION         = os.environ.get("AWS_REGION", "us-east-1")
MODEL_ID       = os.environ.get("MODEL_ID", "us.anthropic.claude-sonnet-4-5-20250929-v1:0")
BUCKET_PREFIX  = "a2a-demo-docs-"
RESULTS_PREFIX = "results/docx/"
MAX_BYTES      = 10 * 1024 * 1024

s3      = boto3.client("s3", region_name=REGION)
bedrock = boto3.client("bedrock-runtime", region_name=REGION)

mcp = FastMCP("aws-docx-agent", host="0.0.0.0", stateless_http=True)

ENRICH_PROMPT = """You are a document analysis agent. Read the document and return ONLY a JSON object with these keys:
"document_type": a short label such as invoice, contract, resume, report or letter,
"key_points": a list of 3 to 6 short strings,
"entities": an object with lists named "people", "organizations", "dates", "amounts",
"summary": a clear summary in 4 to 6 sentences.
Return no markdown and no extra text.

DOCUMENT:
{text}"""


def check_bucket(bucket):
    if not bucket.startswith(BUCKET_PREFIX):
        raise ValueError("Bucket not allowed")


def extract_docx_text(bucket, key):
    check_bucket(bucket)
    obj = s3.get_object(Bucket=bucket, Key=key)
    if obj["ContentLength"] > MAX_BYTES:
        raise ValueError("File is larger than 10 MB")
    doc = Document(io.BytesIO(obj["Body"].read()))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def enrich_and_summarize(text):
    resp = bedrock.converse(
        modelId=MODEL_ID,
        messages=[{"role": "user",
                   "content": [{"text": ENRICH_PROMPT.format(text=text[:30000])}]}],
        inferenceConfig={"maxTokens": 1500})
    raw = resp["output"]["message"]["content"][0]["text"].strip()
    raw = raw.removeprefix("```json").removeprefix("```").removesuffix("```").strip()
    return json.loads(raw)


@mcp.tool()
def process_docx_document(job_id: str, file_name: str, s3_bucket: str, s3_key: str) -> dict:
    """Process a Word (.docx) document on AWS: read it from S3, extract paragraphs and
    tables, enrich it (document type, key points, entities), summarize it with Claude
    on Amazon Bedrock, and store the result in S3. Returns the result."""
    text = extract_docx_text(s3_bucket, s3_key)
    if len(text) < 20:
        raise ValueError("Document has no readable text")
    ai = enrich_and_summarize(text)
    result = {
        "job_id": job_id, "status": "done", "file_name": file_name,
        "document_type": ai.get("document_type", ""),
        "key_points": ai.get("key_points", []),
        "entities": ai.get("entities", {}),
        "summary": ai.get("summary", ""),
        "word_count": len(text.split()),
        "stored_in": f"s3://{s3_bucket}/{RESULTS_PREFIX}{job_id}.json",
        "processed_by": "aws-docx-agent",
    }
    s3.put_object(Bucket=s3_bucket, Key=f"{RESULTS_PREFIX}{job_id}.json",
                  Body=json.dumps(result).encode(), ContentType="application/json")
    return result


@mcp.tool()
def get_docx_result(job_id: str, s3_bucket: str) -> dict:
    """Look up the stored processing result for a DOCX job ID."""
    check_bucket(s3_bucket)
    try:
        obj = s3.get_object(Bucket=s3_bucket, Key=f"{RESULTS_PREFIX}{job_id}.json")
        return json.loads(obj["Body"].read())
    except s3.exceptions.NoSuchKey:
        return {"job_id": job_id, "status": "not_found"}


if __name__ == "__main__":
    mcp.run(transport="streamable-http")
